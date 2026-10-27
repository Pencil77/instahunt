import pytest
from pathlib import Path
import openpyxl

from instahunt.db import init_db, upsert_profile
from instahunt.exporter import SheetExporter, extract_username_from_url


def test_extract_username_from_url():
    assert extract_username_from_url("https://www.instagram.com/sonali.rao22?stkn=MWtkYTJ3ZTFsZHhmMg==") == "sonali.rao22"
    assert extract_username_from_url("https://www.instagram.com/rekhasinghlife/reels/") == "rekhasinghlife"
    assert extract_username_from_url("https://www.instagram.com/reel/DbD8M1JM4N2/?stkn=MTdocnRrcmxtdGF2bw==") is None


def test_import_and_export_sheet(tmp_path):
    # Setup temporary test workbook
    sheet_file = tmp_path / "TestStormShadow.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["IG LINK"])
    ws1.append(["https://www.instagram.com/existing_creator?stkn=abc123xyz"])

    ws2 = wb.create_sheet(title="Example Profiles")
    ws2.append(["Profiles that work"])
    ws2.append(["https://www.instagram.com/sonali.rao22?stkn=MWtkYTJ3ZTFsZHhmMg=="])
    ws2.append(["Profiles that DONT work"])
    ws2.append(["https://www.instagram.com/shraboni_454/reels/"])
    wb.save(sheet_file)

    # Initialize DB in tmp_path
    db_file = tmp_path / "instahunt.db"
    init_db(db_file)

    exporter = SheetExporter(sheet_path=sheet_file, db_path=db_file)
    import_stats = exporter.import_seed_profiles()
    assert import_stats["sheet1_imported"] == 1
    assert import_stats["examples_working_imported"] == 1
    assert import_stats["examples_nonworking_imported"] == 1

    # Add a new qualified profile
    upsert_profile({
        "username": "new_jaipur_creator",
        "full_name": "New Creator",
        "avg_comments": 150.0,
        "detected_city": "Jaipur",
        "city_tier": "Tier 2",
        "status": "qualified",
        "evidence": "Verified test profile",
    }, db_path=db_file)

    # Export to sheet
    output_sheet = tmp_path / "OutputStormShadow.xlsx"
    wb.save(output_sheet)
    # Patch exporter sheet_path
    exporter.sheet_path = output_sheet

    added = exporter.append_qualified_to_sheet()
    assert added == 1

    # Verify output sheet has new row
    wb_out = openpyxl.load_workbook(output_sheet)
    ws_out = wb_out["Sheet1"]
    rows = list(ws_out.iter_rows(values_only=True))
    assert any("new_jaipur_creator" in str(r[0]) for r in rows)
